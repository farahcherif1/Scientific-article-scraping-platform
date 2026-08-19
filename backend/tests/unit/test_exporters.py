import csv
import io
import json
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.db.models import Article, CollectionRun
from app.exporters.csv_exporter import export_articles_to_csv
from app.exporters.graph_exporter import export_articles_to_graph_json
from app.exporters.json_exporter import export_articles_to_json
from app.exporters.xlsx_exporter import export_dataset_to_xlsx
from app.use_cases.export_dataset import ExportDataset, export_params_record
from app.use_cases.list_articles import ArticleFilters, compute_stats


def _article(**overrides) -> Article:
    defaults = {
        "id": 1,
        "collection_run_id": 1,
        "title": "Deep Learning, Meet Genomics",
        "authors": ["Ada Lovelace", "Grace Hopper"],
        "year": 2021,
        "abstract": "An abstract.",
        "url": "https://example.org/a",
        "doi": "10.1/abc",
        "venue": "Venue",
        "domain": "cs",
        "categories": ["cs.AI", "cs.LG"],
        "citation_count": 5,
        "keywords": ["deep learning"],
        "keywords_auto": [],
        "source": "arxiv",
        "search_keyword": "deep learning",
        "collection_date": datetime(2026, 1, 1, tzinfo=UTC),
        "duplicate_group_id": None,
        "duplicate_similarity_score": None,
        "duplicate_rule": None,
        "is_duplicate": False,
        "missing_fields": [],
        "relevance_score": 3,
    }
    defaults.update(overrides)
    return Article(**defaults)


def _dataset(articles: list[Article]) -> ExportDataset:
    run = CollectionRun(id=1, keywords=["deep learning"], sources=["arxiv"], status="completed")
    run.created_at = datetime(2026, 1, 1, tzinfo=UTC)
    deduped = [a for a in articles if not a.is_duplicate]
    duplicates = [a for a in articles if a.is_duplicate]
    return ExportDataset(
        run=run,
        sort="-relevance",
        filters=ArticleFilters(sources=["arxiv"]),
        articles=articles,
        deduped=deduped,
        duplicates=duplicates,
        stats=compute_stats(articles),
    )


class TestCsvExporter:
    def test_output_is_utf8_with_bom(self):
        content = export_articles_to_csv([_article()])
        assert content.startswith(b"\xef\xbb\xbf")

    def test_comma_separated_and_parses_back_with_stdlib_csv(self):
        content = export_articles_to_csv([_article(title="A, Comma Title")])
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        assert reader.fieldnames is not None and "," not in "".join(reader.fieldnames)
        assert len(rows) == 1
        assert rows[0]["title"] == "A, Comma Title"

    def test_list_fields_are_joined_with_semicolons(self):
        content = export_articles_to_csv([_article(authors=["Ada Lovelace", "Grace Hopper"])])
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        assert rows[0]["authors"] == "Ada Lovelace; Grace Hopper"

    def test_missing_values_become_empty_cells_not_the_word_none(self):
        content = export_articles_to_csv([_article(doi=None, year=None)])
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        assert rows[0]["doi"] == ""
        assert rows[0]["year"] == ""

    def test_empty_article_list_still_has_a_header_row(self):
        content = export_articles_to_csv([])
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        assert rows == []


class TestJsonExporter:
    def test_parses_to_a_list_of_records(self):
        content = export_articles_to_json([_article(), _article(id=2)])
        records = json.loads(content)
        assert isinstance(records, list)
        assert len(records) == 2

    def test_record_matches_the_article_schema(self):
        content = export_articles_to_json([_article()])
        record = json.loads(content)[0]
        assert record["title"] == "Deep Learning, Meet Genomics"
        assert record["authors"] == ["Ada Lovelace", "Grace Hopper"]
        assert record["categories"] == ["cs.AI", "cs.LG"]
        assert record["doi"] == "10.1/abc"
        assert record["is_duplicate"] is False
        assert record["collection_date"] == "2026-01-01T00:00:00+00:00"

    def test_ai_ready_fields_are_included_for_export(self):
        content = export_articles_to_json([_article()])
        record = json.loads(content)[0]
        assert record["text_unified"] == "Deep Learning, Meet Genomics An abstract. deep learning"
        assert record["abstract_missing"] is False

    def test_missing_abstract_is_flagged_for_export(self):
        content = export_articles_to_json([_article(abstract=None)])
        record = json.loads(content)[0]
        assert record["text_unified"] == "Deep Learning, Meet Genomics deep learning"
        assert record["abstract_missing"] is True

    def test_missing_fields_stay_null_not_dropped(self):
        content = export_articles_to_json([_article(doi=None)])
        record = json.loads(content)[0]
        assert "doi" in record
        assert record["doi"] is None


class TestGraphJsonExporter:
    def test_creates_article_author_keyword_and_year_nodes_with_links(self):
        content = export_articles_to_graph_json([_article()], "COL-0001")
        graph = json.loads(content)

        assert graph["format"] == "knowledge-graph"
        assert graph["collection_id"] == "COL-0001"
        assert {node["type"] for node in graph["nodes"]} == {"article", "author", "keyword", "year"}
        assert {link["type"] for link in graph["links"]} == {"AUTHORED_BY", "HAS_KEYWORD", "PUBLISHED_IN"}

    def test_reuses_shared_entities_and_ignores_missing_year(self):
        first = _article(id=1, authors=["Ada Lovelace"], keywords=["AI"], year=2021)
        second = _article(
            id=2, authors=["ada lovelace"], keywords=[], keywords_auto=["ai"], year=None
        )
        graph = json.loads(export_articles_to_graph_json([first, second], "COL-0001"))

        assert len([node for node in graph["nodes"] if node["type"] == "author"]) == 1
        assert len([node for node in graph["nodes"] if node["type"] == "keyword"]) == 1
        assert len([node for node in graph["nodes"] if node["type"] == "year"]) == 1


class TestXlsxExporter:
    def test_has_the_five_required_sheets(self):
        content = export_dataset_to_xlsx(_dataset([_article()]))
        workbook = load_workbook(io.BytesIO(content))
        assert workbook.sheetnames == ["articles", "deduped", "duplicates", "stats", "params"]

    def test_header_rows_are_bold_and_frozen_on_every_sheet(self):
        content = export_dataset_to_xlsx(_dataset([_article()]))
        workbook = load_workbook(io.BytesIO(content))
        for name in workbook.sheetnames:
            ws = workbook[name]
            assert ws.freeze_panes == "A2"
            header_cells = [cell for cell in ws[1] if cell.value is not None]
            assert header_cells, f"{name} sheet has no header cells"
            assert all(cell.font.bold for cell in header_cells)

    def test_articles_sheet_contains_all_rows_deduped_only_non_duplicates(self):
        canonical = _article(id=1, is_duplicate=False)
        duplicate = _article(id=2, is_duplicate=True, duplicate_group_id="DUP-001")
        content = export_dataset_to_xlsx(_dataset([canonical, duplicate]))
        workbook = load_workbook(io.BytesIO(content))

        assert workbook["articles"].max_row == 3  # header + 2 rows
        assert workbook["deduped"].max_row == 2  # header + canonical only
        assert workbook["duplicates"].max_row == 2  # header + duplicate only

    def test_stats_sheet_reports_totals(self):
        canonical = _article(id=1, is_duplicate=False)
        duplicate = _article(id=2, is_duplicate=True)
        content = export_dataset_to_xlsx(_dataset([canonical, duplicate]))
        ws = load_workbook(io.BytesIO(content))["stats"]
        values = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2, max_row=5)}
        assert values["Total Articles"] == 2
        assert values["Deduped"] == 1
        assert values["Duplicates"] == 1

    def test_params_sheet_documents_the_run_and_active_filters(self):
        content = export_dataset_to_xlsx(_dataset([_article()]))
        ws = load_workbook(io.BytesIO(content))["params"]
        values = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
        assert values["collection_id"] == "COL-0001"
        assert values["filter_source"] == "arxiv"
        assert values["sort"] == "-relevance"


def test_export_params_record_matches_run_and_filters():
    dataset = _dataset([_article()])
    record = export_params_record(dataset)
    assert record["collection_id"] == "COL-0001"
    assert record["keywords"] == "deep learning"
    assert record["sources"] == "arxiv"
    assert record["filter_source"] == "arxiv"
