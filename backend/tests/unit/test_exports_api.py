import csv
import io
import json
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.models import Article, Base, CollectionRun
from app.db.session import get_db
from app.main import app

engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)


def _override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


client = TestClient(app)


@pytest.fixture(autouse=True)
def _reset_db():
    # Scoped per-test (not a module-level assignment) so this doesn't clobber
    # other test modules' own `app.dependency_overrides[get_db]` when several
    # are collected in the same pytest session (see test_start_collection.py).
    previous_override = app.dependency_overrides.get(get_db)
    Base.metadata.create_all(bind=engine)
    app.dependency_overrides[get_db] = _override_get_db
    try:
        yield
    finally:
        if previous_override is not None:
            app.dependency_overrides[get_db] = previous_override
        else:
            app.dependency_overrides.pop(get_db, None)
        Base.metadata.drop_all(bind=engine)


def _make_article(db, run_id: int, **overrides) -> Article:
    defaults = {
        "collection_run_id": run_id,
        "title": "Sample Article",
        "authors": ["Jane Doe"],
        "year": 2020,
        "abstract": "An abstract.",
        "url": "https://example.org/a",
        "doi": "10.1/abc",
        "venue": "Venue",
        "domain": "cs",
        "categories": ["cs.AI"],
        "citation_count": 5,
        "source": "arxiv",
        "search_keyword": "ai",
        "collection_date": datetime.now(UTC),
        "duplicate_group_id": None,
        "duplicate_similarity_score": None,
        "duplicate_rule": None,
        "is_duplicate": False,
        "missing_fields": [],
        "relevance_score": 1,
    }
    defaults.update(overrides)
    article = Article(**defaults)
    db.add(article)
    db.commit()
    return article


def _seed_run() -> str:
    db = TestingSessionLocal()
    run = CollectionRun(keywords=["ai"], sources=["arxiv", "openalex"], status="completed")
    db.add(run)
    db.commit()
    db.refresh(run)

    _make_article(db, run.id, title="Arxiv Paper", source="arxiv")
    _make_article(db, run.id, title="OpenAlex Paper", source="openalex")
    _make_article(
        db, run.id, title="Duplicate Paper", source="arxiv", is_duplicate=True,
        duplicate_group_id="DUP-001",
    )

    public_id = run.public_id
    db.close()
    return public_id


def test_export_unknown_collection_returns_404():
    response = client.get("/api/v1/collections/COL-9999/export?format=xlsx")
    assert response.status_code == 404


def test_export_requires_format_param():
    run_id = _seed_run()
    response = client.get(f"/api/v1/collections/{run_id}/export")
    assert response.status_code == 422


def test_export_rejects_unknown_format():
    run_id = _seed_run()
    response = client.get(f"/api/v1/collections/{run_id}/export?format=pdf")
    assert response.status_code == 422


def test_export_invalid_year_range_returns_422():
    run_id = _seed_run()
    response = client.get(
        f"/api/v1/collections/{run_id}/export?format=csv&year_from=2022&year_to=2018"
    )
    assert response.status_code == 422


class TestXlsxExport:
    def test_returns_xlsx_content_type_and_attachment_header(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=xlsx")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in response.headers["content-disposition"]
        assert f"{run_id}_export.xlsx" in response.headers["content-disposition"]

    def test_contains_all_five_sheets_with_expected_row_counts(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=xlsx")
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ["articles", "deduped", "duplicates", "stats", "params"]
        assert workbook["articles"].max_row == 4  # header + 3 rows (incl. duplicate)
        assert workbook["deduped"].max_row == 3  # header + 2 non-duplicates
        assert workbook["duplicates"].max_row == 2  # header + 1 duplicate

    def test_active_source_filter_is_reflected_in_the_export(self):
        run_id = _seed_run()
        response = client.get(
            f"/api/v1/collections/{run_id}/export?format=xlsx&source=openalex"
        )
        workbook = load_workbook(io.BytesIO(response.content))
        # Only the OpenAlex row should survive the filter across every sheet.
        assert workbook["articles"].max_row == 2  # header + 1
        titles = [row[1].value for row in workbook["articles"].iter_rows(min_row=2)]
        assert titles == ["OpenAlex Paper"]


class TestCsvExport:
    def test_returns_csv_with_utf8_bom(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=csv")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert response.content.startswith(b"\xef\xbb\xbf")

    def test_active_filter_is_reflected(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=csv&source=arxiv")
        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert {row["title"] for row in rows} == {"Arxiv Paper", "Duplicate Paper"}


class TestJsonExport:
    def test_returns_a_list_of_records(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=json")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/json")
        records = json.loads(response.content)
        assert len(records) == 3
        assert {r["title"] for r in records} == {"Arxiv Paper", "OpenAlex Paper", "Duplicate Paper"}

    def test_active_filter_is_reflected(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=json&source=openalex")
        records = json.loads(response.content)
        assert len(records) == 1
        assert records[0]["title"] == "OpenAlex Paper"


class TestGraphExport:
    def test_returns_graph_json_with_entities_and_links(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=graph")
        graph = json.loads(response.content)

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/json")
        assert f"{run_id}_knowledge_graph.json" in response.headers["content-disposition"]
        assert {node["type"] for node in graph["nodes"]} >= {"article", "author", "year"}
        assert {link["type"] for link in graph["links"]} >= {"AUTHORED_BY", "PUBLISHED_IN"}

    def test_active_filters_are_reflected(self):
        run_id = _seed_run()
        response = client.get(f"/api/v1/collections/{run_id}/export?format=graph&source=openalex")
        graph = json.loads(response.content)

        assert [node["label"] for node in graph["nodes"] if node["type"] == "article"] == ["OpenAlex Paper"]
