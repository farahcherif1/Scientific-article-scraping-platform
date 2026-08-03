"""
Excel export (US-06.1). Five sheets - articles, deduped, duplicates, stats,
params - built with openpyxl. Every sheet's header row is bold and frozen
(`freeze_panes="A2"`) so it stays visible while scrolling a long article list.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.exporters.common import ARTICLE_FIELDS, article_to_flat_row
from app.use_cases.export_dataset import ExportDataset, export_params_record

_HEADER_FONT = Font(bold=True)


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"


def _write_articles_sheet(ws: Worksheet, articles: list) -> None:
    _write_header(ws, list(ARTICLE_FIELDS))
    for article in articles:
        row = article_to_flat_row(article)
        ws.append([row[field] for field in ARTICLE_FIELDS])


def _write_stats_sheet(ws: Worksheet, dataset: ExportDataset) -> None:
    stats = dataset.stats
    _write_header(ws, ["Metric", "Value"])
    ws.append(["Total Articles", stats.total])
    ws.append(["Deduped", stats.deduped])
    ws.append(["Duplicates", stats.duplicates])
    ws.append(["% With DOI", stats.pct_with_doi])
    ws.append(["% With Abstract", stats.pct_with_abstract])

    ws.append([])
    source_header_row = ws.max_row + 1
    ws.append(["Source", "Articles", "% With DOI", "% With Abstract", "% With Year"])
    for cell in ws[source_header_row]:
        cell.font = _HEADER_FONT
    for source_stat in stats.per_source:
        ws.append(
            [
                source_stat.source,
                source_stat.count,
                source_stat.pct_with_doi,
                source_stat.pct_with_abstract,
                source_stat.pct_with_year,
            ]
        )

    ws.append([])
    year_header_row = ws.max_row + 1
    ws.append(["Year", "Articles"])
    for cell in ws[year_header_row]:
        cell.font = _HEADER_FONT
    for year, count in stats.articles_per_year:
        ws.append([year, count])


def _write_params_sheet(ws: Worksheet, dataset: ExportDataset) -> None:
    record = export_params_record(dataset)
    _write_header(ws, ["Field", "Value"])
    for field, value in record.items():
        ws.append([field, value])


def export_dataset_to_xlsx(dataset: ExportDataset) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_articles_sheet(workbook.create_sheet("articles"), dataset.articles)
    _write_articles_sheet(workbook.create_sheet("deduped"), dataset.deduped)
    _write_articles_sheet(workbook.create_sheet("duplicates"), dataset.duplicates)
    _write_stats_sheet(workbook.create_sheet("stats"), dataset)
    _write_params_sheet(workbook.create_sheet("params"), dataset)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
